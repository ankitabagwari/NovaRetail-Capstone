from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt

from analyze_data import (
    load_data,
    calculate_analysis,
)


BASE_DIR = Path(__file__).resolve().parent.parent

TEMPLATE_PATH = (
    BASE_DIR / "assets" / "novaretail_template.pptx"
)

OUTPUT_PATH = (
    BASE_DIR / "output" / "NovaRetail_QBR_Final.pptx"
)

CHART_DIR = BASE_DIR / "charts"
ICON_DIR = BASE_DIR / "assets" / "icons"


# Replace every value below using template_shape_report.txt.
SHAPE_MAP = {
    "slide_1": {
        "period": "subtitle_1",
        "kpi_1_value": "kpi1_value",
        "kpi_1_label": "kpi1_label",
        "kpi_2_value": "kpi2_value",
        "kpi_2_label": "kpi2_label",
        "kpi_3_value": "kpi3_value",
        "kpi_3_label": "kpi3_label",
        "summary": "summary_text",
        "icon": "icon_1"
    },

    "slide_2": {
        "headline": "title_2",
        "chart": "chart_performance",
        "insights": "insights_performance",
        "icon": "icon_2"
    },

    "slide_3": {
        "headline": "title_3",
        "chart": "chart_customer",
        "insights": "insights_customer",
        "icon": "icon_3"
    },

    "slide_4": {
        "headline": "title_4",
        "chart": "chart_risk",
        "insights": "insights_risk",
        "icon": "icon_4"
    },

    "slide_5": {
        "actions": "recommendations_text",
        "closing": "closing_statement",
        "icon": "icon_5"
    }
}


def get_shape(slide, shape_name):
    for shape in slide.shapes:
        if shape.name == shape_name:
            return shape

    available_names = [
        shape.name
        for shape in slide.shapes
    ]

    raise ValueError(
        f"Shape {shape_name!r} was not found.\n"
        f"Available shapes: {available_names}"
    )


def remove_shape(shape):
    shape_element = shape._element
    shape_element.getparent().remove(shape_element)


def replace_text(
    slide,
    shape_name,
    text,
    font_size=None,
    bold=None,
    bullet=False,
):
    shape = get_shape(
        slide,
        shape_name,
    )

    if not shape.has_text_frame:
        raise TypeError(
            f"{shape_name!r} does not have a text frame."
        )

    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    lines = [
        line.strip()
        for line in text.split("\n")
        if line.strip()
    ]

    for index, line in enumerate(lines):
        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()

        cleaned_line = line

        if cleaned_line.startswith("•"):
            cleaned_line = cleaned_line[1:].strip()

        paragraph.text = cleaned_line
        paragraph.alignment = PP_ALIGN.LEFT

        if bullet:
            paragraph.level = 0

        paragraph.space_after = Pt(5)

        for run in paragraph.runs:
            if font_size:
                run.font.size = Pt(font_size)

            if bold is not None:
                run.font.bold = bold


def replace_picture(
    slide,
    placeholder_name,
    picture_path,
):
    placeholder = get_shape(
        slide,
        placeholder_name,
    )

    left = placeholder.left
    top = placeholder.top
    width = placeholder.width
    height = placeholder.height

    remove_shape(placeholder)

    picture = slide.shapes.add_picture(
        str(picture_path),
        left,
        top,
        width,
        height,
    )

    picture.name = (
        f"generated_{picture_path.stem}"
    )


def validate_paths():
    required_paths = [
        TEMPLATE_PATH,
        BASE_DIR
        / "assets"
        / "novaretail_dataset.xlsx",
        CHART_DIR / "revenue_by_region.png",
        CHART_DIR / "nps_by_region.png",
        CHART_DIR / "stockout_rate.png",
        ICON_DIR / "globe_region.png",
        ICON_DIR / "growth_chart.png",
        ICON_DIR / "customer_star.png",
        ICON_DIR / "risk_alert.png",
        ICON_DIR / "lightbulb_action.png",
    ]

    missing = [
        path
        for path in required_paths
        if not path.exists()
    ]

    if missing:
        formatted = "\n".join(
            str(path)
            for path in missing
        )

        raise FileNotFoundError(
            f"Required files are missing:\n{formatted}"
        )


def ensure_mapping_updated():
    unresolved = []

    for slide_key, mappings in SHAPE_MAP.items():
        for purpose, shape_name in mappings.items():
            if shape_name == "REPLACE_WITH_SHAPE_NAME":
                unresolved.append(
                    f"{slide_key}.{purpose}"
                )

    if unresolved:
        raise ValueError(
            "Update SHAPE_MAP with actual names from "
            "template_shape_report.txt:\n"
            + "\n".join(unresolved)
        )


def generate_presentation():
    validate_paths()
    ensure_mapping_updated()

    data = load_data()

    results, _, _, _ = calculate_analysis(data)

    presentation = Presentation(
        TEMPLATE_PATH
    )

    if len(presentation.slides) != 5:
        raise ValueError(
            "The template must contain exactly 5 slides."
        )

    # ---------------------------------
    # Slide 1: Executive summary
    # ---------------------------------

    slide = presentation.slides[0]
    mapping = SHAPE_MAP["slide_1"]

    replace_text(
        slide,
        mapping["period"],
        "Q2 2026 — Quarterly Business Review",
        font_size=16,
    )

    replace_text(
        slide,
        mapping["kpi_1_value"],
        f"${results['total_revenue_usd_m'] / 1000:.2f}B",
        font_size=25,
        bold=True,
    )

    replace_text(
        slide,
        mapping["kpi_1_label"],
        "Total Revenue",
        font_size=11,
    )

    replace_text(
        slide,
        mapping["kpi_2_value"],
        f"+{results['qoq_growth_pct']:.1f}%",
        font_size=25,
        bold=True,
    )

    replace_text(
        slide,
        mapping["kpi_2_label"],
        "QoQ Revenue Growth",
        font_size=11,
    )

    replace_text(
        slide,
        mapping["kpi_3_value"],
        f"{results['weighted_nps']:.1f}",
        font_size=25,
        bold=True,
    )

    replace_text(
        slide,
        mapping["kpi_3_label"],
        "Weighted Customer NPS",
        font_size=11,
    )

    slide_1_summary = (
        "NovaRetail generated $1.28B in Q2 revenue, "
        "recovering 4.5% from the previous quarter. "
        "APAC delivered the strongest quarterly revenue "
        "increase and continues to lead customer "
        "experience. However, APAC's high warehouse "
        "utilization may constrain future growth. "
        "Latin America remains the most urgent concern, "
        "combining declining revenue with deteriorating "
        "customer satisfaction and operational risk."
    )

    replace_text(
        slide,
        mapping["summary"],
        slide_1_summary,
        font_size=13,
    )

    replace_picture(
        slide,
        mapping["icon"],
        ICON_DIR / "globe_region.png",
    )

    # ---------------------------------
    # Slide 2: Business performance
    # ---------------------------------

    slide = presentation.slides[1]
    mapping = SHAPE_MAP["slide_2"]

    replace_text(
        slide,
        mapping["headline"],
        (
            "APAC Drives Growth While Latin America "
            "Continues to Decline"
        ),
        font_size=22,
        bold=True,
    )

    slide_2_insights = (
        "• APAC revenue increased to $330M, the strongest "
        "quarterly increase across the four regions.\n"
        "• North America remained the largest contributor, "
        "reaching $465M in Q2 revenue.\n"
        "• Europe recovered to $345M after its Q1 decline.\n"
        "• Latin America fell to $140M as returns increased "
        "to 8.2%."
    )

    replace_text(
        slide,
        mapping["insights"],
        slide_2_insights,
        font_size=12,
        bullet=True,
    )

    replace_picture(
        slide,
        mapping["chart"],
        CHART_DIR / "revenue_by_region.png",
    )

    replace_picture(
        slide,
        mapping["icon"],
        ICON_DIR / "growth_chart.png",
    )

    # ---------------------------------
    # Slide 3: Customer experience
    # ---------------------------------

    slide = presentation.slides[2]
    mapping = SHAPE_MAP["slide_3"]

    replace_text(
        slide,
        mapping["headline"],
        (
            "Customer Experience Improves Everywhere "
            "Except Latin America"
        ),
        font_size=22,
        bold=True,
    )

    slide_3_insights = (
        "• APAC leads with 62 NPS, 92% CSAT and "
        "2.2-day delivery.\n"
        "• North America increased to 55 NPS while "
        "reducing delivery time to 2.9 days.\n"
        "• Europe reached 53 NPS and 88% CSAT.\n"
        "• Latin America declined to 29 NPS as delivery "
        "time and support demand increased."
    )

    replace_text(
        slide,
        mapping["insights"],
        slide_3_insights,
        font_size=12,
        bullet=True,
    )

    replace_picture(
        slide,
        mapping["chart"],
        CHART_DIR / "nps_by_region.png",
    )

    replace_picture(
        slide,
        mapping["icon"],
        ICON_DIR / "customer_star.png",
    )

    # ---------------------------------
    # Slide 4: Operational risk
    # ---------------------------------

    slide = presentation.slides[3]
    mapping = SHAPE_MAP["slide_4"]

    replace_text(
        slide,
        mapping["headline"],
        (
            "Latin America Faces Critical Risk as "
            "APAC Capacity Tightens"
        ),
        font_size=22,
        bold=True,
    )

    slide_4_insights = (
        "• Latin America reached a 12.5% stockout rate, "
        "the highest across all regions.\n"
        "• Fraud incidents in Latin America increased "
        "to 30 in Q2.\n"
        "• Return processing in Latin America reached "
        "10 days.\n"
        "• APAC warehouse utilization reached 92%, "
        "creating a capacity constraint."
    )

    replace_text(
        slide,
        mapping["insights"],
        slide_4_insights,
        font_size=12,
        bullet=True,
    )

    replace_picture(
        slide,
        mapping["chart"],
        CHART_DIR / "stockout_rate.png",
    )

    replace_picture(
        slide,
        mapping["icon"],
        ICON_DIR / "risk_alert.png",
    )

    # ---------------------------------
    # Slide 5: Recommended actions
    # ---------------------------------

    slide = presentation.slides[4]
    mapping = SHAPE_MAP["slide_5"]

    slide_5_actions = (
        "• Stabilize Latin America inventory and delivery "
        "performance through an immediate recovery plan.\n"
        "• Strengthen fraud monitoring and exception "
        "controls in Latin America.\n"
        "• Expand or rebalance APAC warehouse capacity "
        "before utilization limits growth.\n"
        "• Protect growth in North America and Europe "
        "while monitoring returns and service quality."
    )

    replace_text(
        slide,
        mapping["actions"],
        slide_5_actions,
        font_size=14,
        bullet=True,
    )

    replace_text(
        slide,
        mapping["closing"],
        (
            "Act now on concentrated regional risks while "
            "protecting APAC's growth momentum."
        ),
        font_size=15,
        bold=True,
    )

    replace_picture(
        slide,
        mapping["icon"],
        ICON_DIR / "lightbulb_action.png",
    )

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    presentation.save(OUTPUT_PATH)

    print("\nPresentation generated successfully.")
    print(f"Output: {OUTPUT_PATH}")


if __name__ == "__main__":
    generate_presentation()