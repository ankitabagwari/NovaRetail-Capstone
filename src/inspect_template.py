from pathlib import Path
from pptx import Presentation


# -----------------------------
# Project paths
# -----------------------------

BASE_DIR = Path(__file__).resolve().parent.parent

TEMPLATE_PATH = (
    BASE_DIR
    / "assets"
    / "novaretail_template.pptx"
)

REPORT_PATH = (
    BASE_DIR
    / "template_shape_report.txt"
)


# -----------------------------
# Inspect template
# -----------------------------

def inspect_template():

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Template not found:\n{TEMPLATE_PATH}"
        )

    prs = Presentation(TEMPLATE_PATH)

    print("=" * 80)
    print("NOVARETAIL POWERPOINT TEMPLATE")
    print("=" * 80)

    print(f"\nTemplate location: {TEMPLATE_PATH}")
    print(f"Number of slides: {len(prs.slides)}")

    report = []

    for slide_number, slide in enumerate(
        prs.slides,
        start=1
    ):

        heading = (
            f"\n========== SLIDE {slide_number} =========="
        )

        print(heading)
        report.append(heading)

        for shape_number, shape in enumerate(
            slide.shapes,
            start=1
        ):

            shape_text = ""

            if hasattr(shape, "text"):
                shape_text = (
                    shape.text
                    .replace("\n", " | ")
                    .strip()
                )

            line = (
                f"{shape_number:02d} | "
                f"Name: {shape.name} | "
                f"Type: {shape.shape_type} | "
                f"Left: {shape.left} | "
                f"Top: {shape.top} | "
                f"Width: {shape.width} | "
                f"Height: {shape.height} | "
                f"Text: {shape_text[:200]}"
            )

            print(line)
            report.append(line)

    REPORT_PATH.write_text(
        "\n".join(report),
        encoding="utf-8"
    )

    print("\n" + "=" * 80)
    print("INSPECTION COMPLETED")
    print("=" * 80)

    print(
        f"\nShape report saved here:\n{REPORT_PATH}"
    )


if __name__ == "__main__":
    inspect_template()